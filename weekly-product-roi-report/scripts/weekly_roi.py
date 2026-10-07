"""Read verified completed daily ROI workbooks; export one offline natural week."""
from __future__ import annotations

import argparse
import ast
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
import hashlib
from html import escape
import json
import math
from pathlib import Path
import re
import sys
import tempfile
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

VERSION = '0.1.0'
SKILL = Path(__file__).resolve().parents[1]
CENT = Decimal('0.01')


class ReportError(ValueError):
    pass


def require(condition, message):
    if not condition:
        raise ReportError(message)


def amount(value, context):
    require(isinstance(value, (int, float, Decimal)) and not isinstance(value, bool),
            f'{context}: 缺少数值或公式缓存，请在 Excel 中重算并另存副本')
    result = Decimal(str(value))
    require(result.is_finite(), f'{context}: 非有限数值')
    return result.quantize(CENT, rounding=ROUND_HALF_UP)


def arithmetic(text):
    """Evaluate only literal addition/subtraction; never eval spreadsheet text."""
    require(len(text) < 10000, '公式过长')
    def visit(node):
        if isinstance(node, ast.Expression):
            return visit(node.body)
        if isinstance(node, ast.Constant) and type(node.value) in (int, float):
            return Decimal(str(node.value))
        if isinstance(node, ast.BinOp) and isinstance(node.op, (ast.Add, ast.Sub)):
            left, right = visit(node.left), visit(node.right)
            return left + right if isinstance(node.op, ast.Add) else left - right
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.USub):
            return -visit(node.operand)
        raise ReportError(f'不支持的公式: {text[:100]}')
    try:
        result = visit(ast.parse(text, mode='eval'))
    except (SyntaxError, RecursionError) as exc:
        raise ReportError('无法解析费用/销售算术公式') from exc
    require(result.is_finite(), '公式结果非有限数值')
    return result.quantize(CENT, rounding=ROUND_HALF_UP)


def ratio(cost, sales):
    return sales / cost if cost else None


def filename_date(path):
    matches = re.findall(r'(?<!\d)(20\d{6})(?!\d)', path.stem)
    require(len(matches) == 1, f'{path.name}: 文件名须包含一个 YYYYMMDD 日期')
    try:
        return datetime.strptime(matches[0], '%Y%m%d').date()
    except ValueError as exc:
        raise ReportError(f'{path.name}: 无效日期') from exc


def select_week(input_dir, requested):
    require(input_dir.is_dir(), '输入目录不存在')
    files = sorted(p for p in input_dir.glob('*.xlsx') if not p.name.startswith('~$'))
    require(files, '输入目录没有 .xlsx 文件（不递归扫描）')
    dated = [(filename_date(p), p) for p in files]
    available = {d for d, _ in dated}
    if requested:
        start = date.fromisoformat(requested)
    else:
        mondays = {d - timedelta(days=d.weekday()) for d in available}
        complete = [m for m in mondays if all(m + timedelta(days=i) in available for i in range(7))]
        require(len(complete) == 1, '须指定 --week-start：完整自然周为零个或多个，不能自动选周')
        start = complete[0]
    require(start.weekday() == 0, '--week-start 必须是周一')
    dates = [start + timedelta(days=i) for i in range(7)]
    selected = []
    for day in dates:
        candidates = [p for d, p in dated if d == day]
        require(len(candidates) == 1, f'{day}: 需要唯一日报，实际 {len(candidates)} 份；缺失不能填零，重复不能相加')
        selected.append(candidates[0])
    return dates, selected, [p.name for d, p in dated if d not in dates]


def discover(workbook):
    found = []
    for sheet in workbook:
        for row in sheet.iter_rows():
            columns = defaultdict(list)
            for cell in row:
                if cell.value is not None:
                    columns[str(cell.value).strip()].append(cell.column)
            if '付费总费用' not in columns or '综合投产比' not in columns:
                continue
            sales = [c for name, cols in columns.items() if '真实销售额' in name for c in cols]
            require(len(sales) == 1 and len(columns['付费总费用']) == 1 and len(columns['综合投产比']) == 1,
                    f'{sheet.title}: 费用/销售/ROI 表头不唯一')
            found.append((sheet, row[0].row, columns['付费总费用'][0], sales[0], columns['综合投产比'][0]))
    require(len(found) == 1, '无法唯一识别综合投产主表；请检查 schema 说明')
    return found[0]


def sku_text(value):
    require(value is not None and not isinstance(value, bool), 'SKU 缺失')
    if isinstance(value, (int, float)):
        require(math.isfinite(value) and value == int(value) and abs(value) < 10**15,
                'SKU 数值非整数或超出 Excel 精确标识范围；须使用文本 SKU')
        return str(int(value))
    require(isinstance(value, str) and value.strip() and not value.startswith('='), '无效 SKU')
    return value.strip()


def read_daily(path, expected_date):
    import openpyxl
    from openpyxl.utils.datetime import from_excel
    try:
        formulas = openpyxl.load_workbook(path, data_only=False)
    except (BadZipFile, ParseError, KeyError, OSError, ValueError) as exc:
        raise ReportError(f'{path.name}: 无法读取工作簿: {exc}') from exc
    cached = None
    try:
        cached = openpyxl.load_workbook(path, data_only=True)
        sheet, header, costcol, salescol, roicol = discover(formulas)
        values = cached[sheet.title]
        product_headers = [c for row in sheet.iter_rows(max_row=header) for c in row
                           if isinstance(c.value, str) and c.value.strip().startswith('产品（')]
        stores = [c for row in sheet.iter_rows(max_row=header) for c in row if c.value == '店铺']
        require(len(product_headers) == len(stores) == 1, '产品/店铺表头不唯一')
        pcol, storecol = product_headers[0].column, stores[0].column
        reported = sheet.cell(product_headers[0].row, costcol).value
        if isinstance(reported, (int, float)):
            reported = from_excel(reported, formulas.epoch)
        if isinstance(reported, datetime):
            reported = reported.date()
        require(reported == expected_date, f'{path.name}: 文件名与表内日期不一致')
        records, store, totalrow = [], None, None
        for r in range(header + 1, sheet.max_row + 1):
            name = sheet.cell(r, pcol).value
            if name is None:
                require(all(sheet.cell(r, c).value is None for c in (costcol, salescol, roicol)),
                        f'{sheet.title}!{r}: 有金额但产品为空')
                continue
            name = str(name).strip()
            if name in ('合计', '总计'):
                totalrow = r
                break
            # Carry store only within a vertical merged range; blank is not a mapping.
            raw_store = sheet.cell(r, storecol).value
            if raw_store is None:
                for merged in sheet.merged_cells.ranges:
                    if merged.min_col == merged.max_col == storecol and merged.min_row <= r <= merged.max_row:
                        raw_store = sheet.cell(merged.min_row, storecol).value
                        break
            require(raw_store is not None and str(raw_store).strip(), f'{name}: 缺少店铺身份')
            store = str(raw_store).strip()
            ctx = f'{path.name}/{sheet.title}/{name}'
            cost, sales = [amount(values.cell(r, c).value, ctx) for c in (costcol, salescol)]
            require(cost >= 0 and sales >= 0, f'{ctx}: 当前模板 profile 不支持负金额，请核对口径')
            costformula = sheet.cell(r, costcol).value
            calculated_cost = arithmetic(costformula[1:]) if isinstance(costformula, str) and costformula.startswith('=') else amount(costformula, ctx)
            require(calculated_cost == cost, f'{ctx}: 费用缓存不一致')
            salesformula = sheet.cell(r, salescol).value
            ref = re.fullmatch(r"='((?:[^']|'')+)'!([A-Z]+)(\d+)", str(salesformula))
            require(ref, f'{ctx}: 不支持的主表销售公式')
            linked_name = ref[1].replace("''", "'")
            require(linked_name in formulas.sheetnames, f'{ctx}: 销售引用工作表缺失')
            linked, linked_values = formulas[linked_name], cached[linked_name]
            rr, target = int(ref[3]), f'{ref[2]}{ref[3]}'
            require(linked.cell(rr, 1).value == name, f'{ctx}: 销售来源产品不一致')
            require(amount(linked_values[target].value, ctx) == sales, f'{ctx}: 销售引用缓存不一致')
            sku_columns = sorted({c.column for row in linked.iter_rows(max_row=rr - 1) for c in row
                                  if isinstance(c.value, str) and 'sku' in c.value.lower()})
            skus = [sku_text(linked.cell(rr, c).value) for c in sku_columns if linked.cell(rr, c).value is not None]
            require(skus and len(skus) == len(set(skus)), f'{ctx}: SKU 缺失或重复')
            skus.sort()
            identity = json.dumps([store, skus], ensure_ascii=False, separators=(',', ':'))
            key = hashlib.sha256(identity.encode('utf-8')).hexdigest()
            real_formula = linked[target].value
            real = re.fullmatch(r'=ROUND\(([A-Z]+\d+)-([A-Z]+\d+),2\)', str(real_formula))
            require(real, f'{ctx}: 不支持的真实销售公式')
            gross_cell, brushing_cell = real.groups()
            gross = amount(linked_values[gross_cell].value, ctx)
            brush_value = linked_values[brushing_cell].value
            # Excel's existing ROUND(gross-blank,2) treats blank as zero. Preserve
            # that observed formula result, never label blank as confirmed no brushing.
            brush = Decimal(0) if brush_value is None else amount(brush_value, ctx)
            require(linked[brushing_cell].data_type != 'f', f'{ctx}: 不支持刷单金额公式，需独立校验')
            require((gross - brush).quantize(CENT, rounding=ROUND_HALF_UP) == sales, f'{ctx}: 真实销售缓存不一致')
            gross_formula = re.fullmatch(r'=ROUND\((.*),2\)', str(linked[gross_cell].value))
            require(gross_formula and arithmetic(gross_formula[1]) == gross, f'{ctx}: 销售总额缓存不一致')
            cost_ref, sales_ref = sheet.cell(r, costcol).coordinate, sheet.cell(r, salescol).coordinate
            require(sheet.cell(r, roicol).value == f'=IF({cost_ref}=0,0,{sales_ref}/{cost_ref})', f'{ctx}: 不支持 ROI 公式')
            expected_roi = float(sales / cost) if cost else 0
            cached_roi = values.cell(r, roicol).value
            require(type(cached_roi) in (int, float) and math.isfinite(cached_roi) and abs(cached_roi - expected_roi) < 1e-9,
                    f'{ctx}: ROI 缓存不一致')
            records.append({'key': key, 'store': store, 'product': name, 'skus': skus,
                            'date': expected_date.isoformat(), 'cost_cents': int(cost * 100),
                            'sales_cents': int(sales * 100), 'roi': ratio(int(cost * 100), int(sales * 100)),
                            'brushing_state': 'blank_in_supplied_report' if brush_value is None else 'numeric_in_supplied_report',
                            'source': path.name, 'sheet': sheet.title, 'row': r,
                            'cost_cell': cost_ref, 'sales_cell': sales_ref,
                            'sales_precedent': f'{linked_name}!{target}'})
        require(records and totalrow, f'{path.name}: 缺少产品或合计行')
        # Product amounts after a total would otherwise be silently omitted.
        for r in range(totalrow + 1, sheet.max_row + 1):
            require(all(sheet.cell(r, c).value is None for c in (costcol, salescol, roicol)),
                    f'{path.name}: 合计行之后仍有金额')
        require(len({r['key'] for r in records}) == len(records), f'{path.name}: 同日重复产品身份')
        totals = {f'{metric}_cents': sum(r[f'{metric}_cents'] for r in records) for metric in ('cost', 'sales')}
        for col, metric in ((costcol, 'cost'), (salescol, 'sales')):
            letter = openpyxl.utils.get_column_letter(col)
            require(sheet.cell(totalrow, col).value == f'=SUM({letter}{header + 1}:{letter}{totalrow - 1})',
                    f'{path.name}: 不支持合计公式')
            require(int(amount(values.cell(totalrow, col).value, path.name) * 100) == totals[f'{metric}_cents'],
                    f'{path.name}: {metric} 合计不一致')
        tc, ts = sheet.cell(totalrow, costcol).coordinate, sheet.cell(totalrow, salescol).coordinate
        require(sheet.cell(totalrow, roicol).value == f'=IF({tc}=0,0,{ts}/{tc})', f'{path.name}: 不支持合计 ROI 公式')
        total_roi = values.cell(totalrow, roicol).value
        expected = ratio(totals['cost_cents'], totals['sales_cents']) or 0
        require(type(total_roi) in (int, float) and math.isfinite(total_roi) and abs(total_roi - expected) < 1e-9,
                f'{path.name}: 合计 ROI 缓存不一致')
        return records, {'date': expected_date.isoformat(), 'products': len(records), **totals,
                         'roi': ratio(totals['cost_cents'], totals['sales_cents'])}, {
                             'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                             'date': expected_date.isoformat()}
    finally:
        formulas.close()
        if cached is not None:
            cached.close()


def analyze(input_dir, week_start=None):
    dates, selected, ignored = select_week(input_dir, week_start)
    records, audits, sources = [], [], []
    names, reverse = {}, {}
    for day, path in zip(dates, selected):
        daily, audit, source = read_daily(path, day)
        for record in daily:
            key = record['key']
            require(key not in names or names[key] == record['product'], f'{day}: 同一 SKU 身份产品名变化，须核对映射')
            names[key] = record['product']
            for sku in record['skus']:
                identity = (record['store'], sku)
                require(identity not in reverse or reverse[identity] == key, f'{day}: SKU 绑定冲突')
                reverse[identity] = key
        records.extend(daily)
        audits.append(audit)
        sources.append(source)
    grouped = defaultdict(list)
    for record in records:
        grouped[record['key']].append(record)
    iso_dates = [d.isoformat() for d in dates]
    products = []
    for key, daily in grouped.items():
        require([r['date'] for r in daily] == iso_dates, f'{names[key]}: 缺少产品日记录；不能推断为零')
        cost = sum(r['cost_cents'] for r in daily)
        sales = sum(r['sales_cents'] for r in daily)
        products.append({'key': key, 'product': names[key], 'store': daily[0]['store'],
                         'skus': daily[0]['skus'], 'cost_cents': cost, 'sales_cents': sales,
                         'roi': ratio(cost, sales), 'daily': daily})
    products.sort(key=lambda p: (-p['cost_cents'], p['key']))
    totals = {f'{metric}_cents': sum(p[f'{metric}_cents'] for p in products) for metric in ('cost', 'sales')}
    require(max(totals.values()) <= 2**53 - 1, '金额超出浏览器整数精确范围')
    totals['roi'] = ratio(totals['cost_cents'], totals['sales_cents'])
    return {'version': VERSION, 'profile': 'completed_daily_roi_v1',
            'period': {'start': iso_dates[0], 'end': iso_dates[-1], 'dates': iso_dates,
                       'kind': 'natural_week_monday_sunday'},
            'products': products, 'totals': totals, 'daily_totals': audits, 'sources': sources,
            'audit': {'status': 'PASS', 'record_count': len(records), 'product_count': len(products),
                      'main_totals_match': True, 'relevant_formulas_verified': True, 'sku_bindings_stable': True,
                      'blank_brushing_days': sum(r['brushing_state'] == 'blank_in_supplied_report' for r in records),
                      'ignored_outside_week_files': ignored,
                      'scope': 'Supplied completed-report values and bounded formulas; no upstream ledger verification.'}}


def render(analysis):
    def day(record):
        return {'cost': record['cost_cents'], 'sales': record['sales_cents'], 'roi': record['roi']}
    name_counts = Counter(p['product'] for p in analysis['products'])
    data = {'dates': analysis['period']['dates'], 'total': day(analysis['totals']),
            'daily': [day(r) for r in analysis['daily_totals']],
            'products': [{'id': p['key'], 'name': p['product'] if name_counts[p['product']] == 1
                          else f"{p['product']}（{p['store']} / {','.join(p['skus'])}）",
                          **day(p), 'days': [day(r) for r in p['daily']]} for p in analysis['products']]}
    payload = json.dumps(data, ensure_ascii=False, separators=(',', ':'), allow_nan=False)
    payload = payload.replace('<', '\\u003c').replace('>', '\\u003e').replace('&', '\\u0026')
    start, end = analysis['period']['start'], analysis['period']['end']
    template = (SKILL / 'assets' / 'report.html').read_text(encoding='utf-8')
    title = f'单品自然周投产报告 {start}—{end}'
    subtitle = f"{start.replace('-', '/')}—{end.replace('-', '/')}（周一至周日） · {len(data['products'])} 个产品 · 同店铺 SKU 组合合并口径"
    replacements = {'__ROI_DATA__': payload, '__TITLE__': escape(title), '__SUBTITLE__': escape(subtitle),
                    '__D3_LICENSE__': escape((SKILL / 'assets' / 'D3-LICENSE.txt').read_text(encoding='utf-8')),
                    '__D3__': (SKILL / 'assets' / 'd3-7.9.0.min.js').read_text(encoding='utf-8').replace('</script', '<\\/script')}
    for placeholder, value in replacements.items():
        require(template.count(placeholder) == 1, f'模板占位符不唯一: {placeholder}')
        template = template.replace(placeholder, value)
    return template


def run(input_dir, output_dir, week_start=None):
    analysis = analyze(input_dir, week_start)
    html = render(analysis)
    # A new run directory is an atomic publication boundary. On any validation
    # failure, no success-looking report is written and existing runs stay intact.
    require(not output_dir.exists(), '输出目录已存在；请使用新的运行目录')
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix='.weekly-roi-', dir=output_dir.parent))
    try:
        filename = f"weekly-roi-{analysis['period']['start']}-{analysis['period']['end']}.html"
        (stage / filename).write_text(html, encoding='utf-8')
        (stage / 'analysis.json').write_text(json.dumps(analysis, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
        stage.rename(output_dir)
    finally:
        if stage.exists():
            for child in stage.iterdir():
                child.unlink()
            stage.rmdir()
    return {'status': 'COMPLETE', 'html': str((output_dir / filename).resolve()),
            'analysis': str((output_dir / 'analysis.json').resolve()),
            'period': analysis['period'], 'audit': analysis['audit']}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--version', action='version', version=VERSION)
    commands = parser.add_subparsers(dest='command', required=True)
    commands.add_parser('preflight')
    report = commands.add_parser('run')
    report.add_argument('--input-dir', type=Path, required=True)
    report.add_argument('--output-dir', type=Path, required=True)
    report.add_argument('--week-start', help='YYYY-MM-DD; Monday')
    args = parser.parse_args()
    try:
        import openpyxl
        if args.command == 'preflight':
            require(all((SKILL / name).is_file() for name in ('assets/report.html', 'assets/d3-7.9.0.min.js', 'assets/D3-LICENSE.txt')), '技能资源不完整')
            result = {'status': 'PASS', 'python': sys.version.split()[0], 'openpyxl': openpyxl.__version__, 'skill_version': VERSION}
        else:
            result = run(args.input_dir.resolve(), args.output_dir.resolve(), args.week_start)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (ReportError, ValueError, OSError, ImportError, KeyError, BadZipFile, ParseError) as exc:
        print(json.dumps({'status': 'FAIL', 'reason': str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    raise SystemExit(main())
